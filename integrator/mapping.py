from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal
from pathlib import Path


@dataclass(frozen=True)
class Product:
    barcode: str
    name: str
    unit: str = "BUC"
    fgo_code: str = ""
    vat: Decimal = Decimal("21")
    fgo_verified: bool = False


def load_mapping(path: str | Path, identifier="barcode") -> dict[str, Product]:
    from openpyxl import load_workbook
    if Path(path).suffix.lower() != ".xlsx":
        raise ValueError("Folosește un fișier .xlsx, cu barcode păstrat ca text.")
    wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
    try:
        sheet = wb["Mapare"] if "Mapare" in wb.sheetnames else wb.worksheets[0]
        if (sheet.max_row or 0) > 100001:
            raise ValueError("Fișierul depășește 100.000 de produse.")
        rows = sheet.iter_rows()
        headers = [str(c.value or "").strip().lower() for c in next(rows)]
        if len([h for h in headers if h]) != len(set(h for h in headers if h)):
            raise ValueError("Antete duplicate în Excel.")
        for column in (identifier,):
            if column not in headers:
                raise ValueError(f"Lipsește coloana {column} din primul rând.")
        if not {"cod_fgo", "denumire_factura"}.intersection(headers):
            raise ValueError("Lipsește coloana cod_fgo. Excelul poate conține doar barcode și cod_fgo.")
        products = {}
        for number, cells in enumerate(rows, 2):
            if number > 100001:
                raise ValueError("Fișierul depășește 100.000 de produse.")
            data = {("barcode" if h == identifier else h): cells[i] for i, h in enumerate(headers) if i < len(cells) and h in {identifier, "denumire_factura", "um", "cod_fgo", "tva"}}
            if all(c.value is None or c.value == "" for c in data.values()):
                continue
            if any(c.data_type == "f" for c in data.values()):
                raise ValueError(f"Rândul {number}: înlocuiește formulele cu valori.")
            raw = data["barcode"].value if "barcode" in data else None
            if not isinstance(raw, str) or not raw.strip():
                raise ValueError(f"Rândul {number}: barcode trebuie introdus ca TEXT, pentru a păstra zerourile inițiale.")
            code = raw.strip()
            if code in products:
                raise ValueError(f"Barcode duplicat la rândul {number}: {code}")
            def val(k, default=""):
                cell = data.get(k)
                return str(cell.value).strip() if cell is not None and cell.value is not None else default
            name, unit = val("denumire_factura"), val("um") or "BUC"
            fgo_code = val("cod_fgo")
            if fgo_code and not isinstance(data["cod_fgo"].value, str):
                raise ValueError(f"Rândul {number}: cod_fgo trebuie introdus ca TEXT, inclusiv zerourile inițiale.")
            if (not name and not fgo_code) or len(name) > 1000 or len(unit) > 5 or len(code) > 128:
                raise ValueError(f"Rândul {number}: completează cod_fgo sau denumire_factura; verifică lungimea codului și a UM.")
            # A mapped FGO code owns the article details; any old Excel details are ignored.
            if fgo_code:
                name, unit = "", ""
            vat = Decimal("21") if fgo_code else Decimal(val("tva") or "21")
            if vat not in (Decimal("21"), Decimal("11")):
                raise ValueError(f"Rândul {number}: TVA acceptat 21 sau 11, exprimat ca număr, fără %.")
            if len(fgo_code) > 128:
                raise ValueError(f"Rândul {number}: cod_fgo depășește 128 de caractere.")
            products[code] = Product(code, name, unit, fgo_code, vat)
        if not products:
            raise ValueError("Excelul nu conține produse. Completează cel puțin un barcode și codul articolului FGO.")
        return products
    finally:
        wb.close()
