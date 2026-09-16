from copy import deepcopy
from dataclasses import replace
from datetime import date
from decimal import Decimal
from unittest import TestCase
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from integrator.config import Settings, load_settings, save_settings
from integrator.domain import build_draft, signature
from integrator.emag import EmagAPI, envelope
from integrator.mapping import Product, load_shared_mapping, shared_rows, write_shared_mapping
from integrator.ui import App
from support import workspace_temp
from test_emag_reports import emag_order


EAN = "0059400000001"


class SharedMappingTests(TestCase):
    def test_roundtrip_preserves_leading_zero_and_platform_namespaces(self):
        with workspace_temp() as folder:
            path = folder / "common.xlsx"
            write_shared_mapping(path, [("SKU-A", EAN, ""), ("SKU-B", "", EAN)])
            maps = load_shared_mapping(path)
            self.assertEqual(maps["trendyol"][EAN].fgo_code, "SKU-A")
            self.assertEqual(maps["emag"][EAN].fgo_code, "SKU-B")
            write_shared_mapping(path, shared_rows(maps["trendyol"], maps["emag"]))
            self.assertEqual(load_shared_mapping(path), maps)

    def test_duplicate_ean_in_same_platform_is_rejected(self):
        with workspace_temp() as folder:
            path = folder / "common.xlsx"
            write_shared_mapping(path, [("SKU-A", EAN, ""), ("SKU-B", EAN, "")])
            with self.assertRaisesRegex(ValueError, "duplicat"):
                load_shared_mapping(path)

    def test_numeric_and_formula_identifiers_are_rejected(self):
        with workspace_temp() as folder:
            path = folder / "common.xlsx"
            for value in (594000000001, '=TEXT(1,"0000000000000")'):
                book = Workbook()
                book.active.append(["cod_fgo", "ean_trendyol", "ean_emag"])
                book.active.append(["SKU", value, ""])
                book.save(path)
                with self.assertRaisesRegex(ValueError, "TEXT"):
                    load_shared_mapping(path)

    def test_export_keeps_article_codes_literal(self):
        with workspace_temp() as folder:
            path = folder / "common.xlsx"
            write_shared_mapping(path, [("=SKU", EAN, "")])
            self.assertEqual(load_shared_mapping(path)["trendyol"][EAN].fgo_code, "=SKU")


class EmagEanTests(TestCase):
    def setUp(self):
        self.settings = replace(Settings(), unified_mapping_path="common.xlsx")
        self.mapping = {EAN: Product(EAN, "Articol FGO", fgo_code="SKU", fgo_verified=True)}

    def test_offer_lookup_matches_id_and_preserves_invoice_fingerprint(self):
        raw = emag_order()
        before = signature(raw)
        api = EmagAPI(self.settings)
        with patch.object(api, "call", return_value={"results": [{"id":123,"ean":[EAN]}]}) as call:
            order = api.enrich_eans(raw["_emag"], "RO")
            self.assertEqual(call.call_args.args, ("RO", "/product_offer/read", {"id":123,"itemsPerPage":100,"currentPage":1}))
            self.assertEqual(signature(raw), before)
            self.assertNotIn("ean", order["products"][0])
            normalized = envelope(order, "RO")
            self.assertEqual(normalized["lines"][0]["barcode"], EAN)
            draft = build_draft(normalized, self.mapping, self.settings)
            self.assertEqual(draft.total, Decimal("36.30"))
            self.assertEqual(draft.payload["Continut"][0]["CodArticol"], "SKU")
            api.enrich_eans(deepcopy(order), "RO")
            self.assertEqual(call.call_count, 1)
            api.enrich_eans(deepcopy(order), "RO", refresh=True)
            self.assertEqual(call.call_count, 2)

    def test_offer_mismatch_or_missing_ean_never_uses_product_id(self):
        api = EmagAPI(self.settings)
        for result in ([{"id":999,"ean":[EAN]}], [{"id":123,"ean":[]}], [{"id":123,"ean":594000000001}]):
            with patch.object(api, "call", return_value={"results": result}), self.assertRaises(ValueError):
                api.enrich_eans(emag_order()["_emag"], "RO", refresh=True)
        with self.assertRaisesRegex(ValueError, "EAN"):
            build_draft(emag_order(), {"123":self.mapping[EAN]}, self.settings)

    def test_ambiguous_eans_block_invoice(self):
        raw = emag_order()
        raw["_emag"]["products"][0]["ean"] = [EAN, "0059400000002"]
        mapping = {**self.mapping, "0059400000002":Product("0059400000002", "Alt produs", fgo_code="OTHER", fgo_verified=True)}
        with self.assertRaisesRegex(ValueError, "diferite"):
            build_draft(raw, mapping, self.settings)

    def test_failed_lookup_remains_visible_with_explanation(self):
        api = EmagAPI(self.settings)
        def response(market, path, payload):
            if path == "/order/read":return {"results":[emag_order()["_emag"]]}
            return {"results":[]}
        with patch.object(api, "call", side_effect=response):
            rows = api.fetch("RO", date.today(), date.today())
        self.assertEqual(len(rows),1)
        with self.assertRaisesRegex(ValueError, "identificat"):
            build_draft(rows[0], self.mapping, self.settings)


class SharedWorkspaceTests(TestCase):
    def test_same_platform_structure_and_combined_pending_routing(self):
        with workspace_temp() as folder:
            app = App(folder)
            try:
                app.withdraw()
                self.assertEqual([app.notebook.tab(t,"text") for t in app.notebook.tabs()], ["Trendyol","eMAG","De facturat","Parametrizare","Rapoarte","Ghid & fiscalitate"])
                for tab in (app.trendyol_tab,app.emag_tab):
                    self.assertEqual([tab.tabs.tab(t,"text") for t in tab.tabs.tabs()], ["Comenzi & facturare","Istoric"])
                app.settings.unified_mapping_path = "common.xlsx"
                app.emag_tab.open_profile()
                app.emag_tab.mapping = {EAN:Product(EAN,"Produs comun",fgo_code="SKU",fgo_verified=True)}
                raw = emag_order()
                raw["_emag"]["products"][0]["ean"] = [EAN]
                app.emag_tab.store.put_orders([raw])
                app.refresh()
                self.assertEqual(app.pending_tree.set("990001:0", "platform"), "Trendyol")
                self.assertEqual(app.pending_tree.set("emag:RO:12345:0", "platform"), "eMAG")
                self.assertEqual(app.pending_tree.set("emag:RO:12345:0", "barcode"), EAN)
                drafts,errors = app.prepare_selected(["990001","emag:RO:12345","990001"])
                self.assertFalse(errors)
                self.assertEqual(len(drafts), 2)
                self.assertEqual(drafts[1].total, Decimal("36.30"))
                self.assertIs(app.service_for_pid("emag:RO:12345").store, app.emag_tab.store)
                self.assertIs(app.service_for_pid("990001").store, app.store)
                app.report_year.set("2025")
                app.report_whole_year()
                self.assertEqual(app.report_start.get(), "2025-01-01")
                self.assertEqual(app.report_end.get(), "2025-12-31")
            finally:
                app.destroy()

    def test_existing_mapping_migrates_without_losing_settings_or_original(self):
        with workspace_temp() as folder:
            old = folder / "old.xlsx"
            book = Workbook()
            book.active.append(["barcode","cod_fgo"])
            book.active.append([EAN,"SKU"])
            book.save(old)
            original = old.read_bytes()
            settings = replace(Settings(), mode="production", cui="10000000", seller_id="123", mapping_path=str(old), api_key="example-api-key")
            save_settings(folder / "settings.json", settings)
            app = App(folder)
            try:
                app.withdraw()
                shared = app.settings.unified_mapping_path
                self.assertTrue(shared)
                self.assertEqual(app.settings.api_key, settings.api_key)
                self.assertEqual(load_shared_mapping(shared)["trendyol"][EAN].fgo_code, "SKU")
                self.assertEqual(old.read_bytes(), original)
                self.assertEqual(load_settings(folder / "settings.json").unified_mapping_path, shared)
            finally:
                app.destroy()
            app = App(folder)
            try:
                app.withdraw()
                self.assertEqual(app.fields["unified_mapping_path"].get(), shared)
                self.assertIn(EAN, app.mapping)
            finally:
                app.destroy()
